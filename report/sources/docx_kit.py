"""Kit mínimo para relatórios académicos, assente em python-docx.

Marcação em linha suportada no texto corrente, nas listas, nas células de
tabela e nas legendas:

    `identificador`   -> tipo monoespaçado, realce discreto, sem verificação
                         ortográfica (os termos de código estão em inglês)
    **texto**         -> negrito

O objetivo da marcação de código é separar visualmente os identificadores do
código-fonte da prosa em português.
"""
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "Calibri"
HEAD_FONT = "Calibri"
CODE_FONT = "Consolas"

INK = RGBColor(0x1A, 0x1A, 0x1A)
SOFT = RGBColor(0x55, 0x55, 0x55)
CODE_INK = RGBColor(0x1B, 0x3A, 0x5C)
RULE = "BFBFBF"
CODE_BG = "F4F4F6"
CODE_CHIP = "EEEFF3"
HDR_BG = "E4E9F2"

DOC_LANG = "pt-PT"
CODE_LANG = "en-GB"

TEXT_WIDTH_CM = 15.5


# --------------------------------------------------------------- baixo nível
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


def _field(paragraph, instr, placeholder=None):
    r = paragraph.add_run()
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = instr
    r._r.append(t)
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    if placeholder:
        ph = OxmlElement("w:t")
        ph.text = placeholder
        r._r.append(ph)
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
    return r


# O esquema OOXML impõe a ordem dos filhos destes contentores e o Word recusa o
# ficheiro se ela for violada. Cada elemento é inserido antes dos que, segundo o
# esquema, lhe sucedem; assim a ordem em que os helpers são chamados deixa de
# importar.
_PPR_ORDER = (
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
    "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
    "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange")

_TCPR_ORDER = (
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
    "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns",
    "cellDel", "cellMerge", "tcPrChange")

_TRPR_ORDER = (
    "cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit",
    "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden", "ins", "del", "trPrChange")

_RPR_ORDER = (
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
    "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish",
    "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight",
    "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
    "eastAsianLayout", "specVanish", "oMath")


def _after(order, tag):
    """Tags que, no esquema, sucedem a `tag` dentro do mesmo contentor."""
    return tuple(f"w:{t}" for t in order[order.index(tag) + 1:])


def _insert_ordered(parent, order, tag, **attrs):
    parent.insert_element_before(_el(f"w:{tag}", **attrs), *_after(order, tag))


_SHD = {"w:val": "clear", "w:color": "auto"}


def _pshade(paragraph, fill):
    _insert_ordered(paragraph._p.get_or_add_pPr(), _PPR_ORDER, "shd",
                    **_SHD, **{"w:fill": fill})


def _cellshade(cell, fill):
    _insert_ordered(cell._tc.get_or_add_tcPr(), _TCPR_ORDER, "shd",
                    **_SHD, **{"w:fill": fill})


def _keep_with_next(paragraph):
    _insert_ordered(paragraph._p.get_or_add_pPr(), _PPR_ORDER, "keepNext",
                    **{"w:val": "1"})


def _no_split(row):
    _insert_ordered(row._tr.get_or_add_trPr(), _TRPR_ORDER, "cantSplit")


def _header_row(row):
    _insert_ordered(row._tr.get_or_add_trPr(), _TRPR_ORDER, "tblHeader")


def _set_lang(rpr, lang, no_proof=False):
    _insert_ordered(rpr, _RPR_ORDER, "lang", **{"w:val": lang})
    if no_proof:
        _insert_ordered(rpr, _RPR_ORDER, "noProof", **{"w:val": "1"})


# ----------------------------------------------------------- marcação inline
_TOKEN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")


def _style_code_run(run, size):
    """Tipo monoespaçado, cor própria e realce de caracter."""
    run.font.name = CODE_FONT
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), CODE_FONT)
    rpr.rFonts.set(qn("w:hAnsi"), CODE_FONT)
    rpr.rFonts.set(qn("w:cs"), CODE_FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = CODE_INK
    _insert_ordered(rpr, _RPR_ORDER, "shd", **_SHD, **{"w:fill": CODE_CHIP})
    # Os identificadores estão em inglês: não devem ser assinalados pelo
    # corretor ortográfico português.
    _set_lang(rpr, CODE_LANG, no_proof=True)


def emit(par, text, bold=False, italic=False, size=None, color=None):
    """Escreve `text` no parágrafo, interpretando `código` e **negrito**."""
    code_size = (size or 11) * 0.87
    for chunk in _TOKEN.split(text):
        if not chunk:
            continue
        if len(chunk) > 2 and chunk[0] == "`" and chunk[-1] == "`":
            run = par.add_run(chunk[1:-1])
            _style_code_run(run, code_size)
            run.bold = bold or None
            continue
        if len(chunk) > 4 and chunk.startswith("**") and chunk.endswith("**"):
            run = par.add_run(chunk[2:-2])
            run.bold = True
            run.italic = italic
        elif len(chunk) > 2 and chunk[0] == "*" and chunk[-1] == "*":
            run = par.add_run(chunk[1:-1])
            run.bold = bold
            run.italic = True
        else:
            run = par.add_run(chunk)
            run.bold = bold
            run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return par


# ------------------------------------------------------------------ documento
class Report:
    def __init__(self):
        self.doc = Document()
        self.fig_n = 0
        self.tab_n = 0
        self.lst_n = 0
        self._setup_page()
        self._setup_styles()
        self._request_field_update()

    # ---- configuração
    def _request_field_update(self):
        """Pede ao leitor que recalcule os campos ao abrir.

        Sem isto o índice e as listas de figuras e tabelas ficam vazios até
        alguém os atualizar à mão.
        """
        settings = self.doc.settings.element
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.append(el)

    def _setup_page(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(3.0)
        s.right_margin = Cm(2.5)
        s.header_distance = s.footer_distance = Cm(1.25)

    def _setup_styles(self):
        st = self.doc.styles

        n = st["Normal"]
        n.font.name = BODY_FONT
        n.font.size = Pt(11)
        n.font.color.rgb = INK
        n.paragraph_format.space_after = Pt(8)
        n.paragraph_format.line_spacing = 1.35
        n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        n._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        _set_lang(n._element.rPr, DOC_LANG)

        for name, size, before, after, color, bold in [
            ("Heading 1", 17, 0, 12, INK, True),
            ("Heading 2", 13.5, 16, 7, INK, True),
            ("Heading 3", 11.5, 12, 5, INK, True),
            ("Heading 4", 11, 10, 4, SOFT, True),
        ]:
            h = st[name]
            h.font.name = HEAD_FONT
            h.font.size = Pt(size)
            h.font.bold = bold
            h.font.color.rgb = color
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.line_spacing = 1.12

        c = st.add_style("BlumeCaption", 1)
        c.base_style = st["Normal"]
        c.font.size = Pt(9)
        c.font.color.rgb = SOFT
        c.paragraph_format.space_before = Pt(4)
        c.paragraph_format.space_after = Pt(12)
        c.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.line_spacing = 1.15

        k = st.add_style("BlumeCode", 1)
        k.base_style = st["Normal"]
        k.font.name = CODE_FONT
        k.font.size = Pt(8.5)
        k._element.rPr.rFonts.set(qn("w:ascii"), CODE_FONT)
        k._element.rPr.rFonts.set(qn("w:hAnsi"), CODE_FONT)
        k.paragraph_format.space_before = Pt(0)
        k.paragraph_format.space_after = Pt(0)
        k.paragraph_format.line_spacing = 1.1
        k.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        k.paragraph_format.left_indent = Cm(0.4)
        _set_lang(k._element.rPr, CODE_LANG, no_proof=True)

        t = st.add_style("BlumeTableText", 1)
        t.base_style = st["Normal"]
        t.font.size = Pt(9)
        t.paragraph_format.space_before = Pt(2)
        t.paragraph_format.space_after = Pt(2)
        t.paragraph_format.line_spacing = 1.12
        t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

        b = st.add_style("BlumeBullet", 1)
        b.base_style = st["Normal"]
        b.paragraph_format.left_indent = Cm(0.75)
        b.paragraph_format.space_after = Pt(4)

        # Estilos latentes: só existem depois de o leitor gerar o índice, pelo
        # que têm de ser criados aqui para não herdarem o entrelinhamento do corpo.
        for name, indent in [("TOC 1", 0), ("TOC 2", 0.5), ("TOC 3", 1.0)]:
            try:
                s = st[name]
            except KeyError:
                s = st.add_style(name, 1)
                s.base_style = st["Normal"]
            s.font.size = Pt(10.5)
            s.paragraph_format.line_spacing = 1.05
            s.paragraph_format.space_before = Pt(0)
            s.paragraph_format.space_after = Pt(3)
            s.paragraph_format.left_indent = Cm(indent)
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

        try:
            tof = st["Table of Figures"]
        except KeyError:
            tof = st.add_style("Table of Figures", 1)
            tof.base_style = st["Normal"]
        tof.font.size = Pt(10.5)
        tof.paragraph_format.line_spacing = 1.05
        tof.paragraph_format.space_before = Pt(0)
        tof.paragraph_format.space_after = Pt(3)
        tof.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # ---- rodapé com numeração
    def add_page_numbers(self):
        for section in self.doc.sections:
            # A capa não é numerada.
            section.different_first_page_header_footer = True
            p = section.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _field(p, " PAGE ", placeholder="1")
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = SOFT
                run.font.name = BODY_FONT

    # ---- estrutura
    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h(self, level, text, in_toc=True):
        par = self.doc.add_heading(text, level=level)
        if not in_toc:
            # Nível 9 é corpo de texto: mantém o aspeto de título mas tira-o do índice.
            _insert_ordered(par._p.get_or_add_pPr(), _PPR_ORDER, "outlineLvl",
                            **{"w:val": "9"})
        return par

    def p(self, text="", style=None, align=None, size=None, italic=False, bold=False,
          space_after=None, color=None):
        par = self.doc.add_paragraph(style=style)
        if text:
            emit(par, text, bold=bold, italic=italic, size=size, color=color)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "left":
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if space_after is not None:
            par.paragraph_format.space_after = Pt(space_after)
        return par

    def bullets(self, items):
        for it in items:
            par = self.doc.add_paragraph(style="BlumeBullet")
            par.paragraph_format.first_line_indent = Cm(-0.35)
            par.paragraph_format.left_indent = Cm(0.75)
            marker = par.add_run("\u2013\u2003")
            marker.font.color.rgb = SOFT
            emit(par, it)

    # ---- legendas com campo SEQ (alimentam as listas de figuras/tabelas)
    def _caption(self, label, number, text, align="center", keep=False):
        cap = self.doc.add_paragraph(style="BlumeCaption")
        if align == "left":
            cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if keep:
            _keep_with_next(cap)
        head = cap.add_run(f"{label} ")
        head.bold = True
        seq = _field(cap, f" SEQ {label} \\* ARABIC ", placeholder=str(number))
        seq.bold = True
        tail = cap.add_run(" \u2014 ")
        tail.bold = True
        emit(cap, text, size=9, color=SOFT)
        return cap

    # ---- figuras / tabelas / listagens
    def figure(self, path, caption, width_cm=15.0):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(8)
        par.paragraph_format.space_after = Pt(2)
        _keep_with_next(par)
        par.add_run().add_picture(path, width=Cm(width_cm))
        self.fig_n += 1
        self._caption("Figura", self.fig_n, caption)
        return self.fig_n

    def _set_widths(self, t, widths):
        """Fixa a largura das colunas.

        O LibreOffice lê `w:tblGrid` e ignora `w:tcW`, pelo que ambos têm de ser
        escritos; as larguras são reescaladas para a mancha de texto.
        """
        scale = TEXT_WIDTH_CM / sum(widths)
        widths = [w * scale for w in widths]

        grid = t._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            for col, w in zip(grid.findall(qn("w:gridCol")), widths):
                col.set(qn("w:w"), str(Cm(w).twips))

        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)

    def table(self, headers, rows, caption, widths=None, font_size=9):
        self.tab_n += 1
        cap = self._caption("Tabela", self.tab_n, caption, align="left", keep=True)
        cap.paragraph_format.space_before = Pt(10)
        cap.paragraph_format.space_after = Pt(4)

        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False

        hdr = t.rows[0]
        _header_row(hdr)
        _no_split(hdr)
        for i, htext in enumerate(headers):
            cell = hdr.cells[i]
            cell.text = ""
            par = cell.paragraphs[0]
            par.style = self.doc.styles["BlumeTableText"]
            emit(par, htext, bold=True, size=font_size)
            _cellshade(cell, HDR_BG)

        for row_data in rows:
            row = t.add_row()
            _no_split(row)
            for i, val in enumerate(row_data):
                cell = row.cells[i]
                cell.text = ""
                par = cell.paragraphs[0]
                par.style = self.doc.styles["BlumeTableText"]
                text, bold = (val if isinstance(val, tuple) else (val, False))
                emit(par, text, bold=bold, size=font_size)

        if widths:
            self._set_widths(t, widths)

        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(10)
        spacer.paragraph_format.space_before = Pt(0)
        return t

    def code(self, lines, caption=None):
        """Bloco de código literal: sem interpretação de marcação."""
        first = None
        for i, line in enumerate(lines):
            par = self.doc.add_paragraph(style="BlumeCode")
            _pshade(par, CODE_BG)
            if i == 0:
                par.paragraph_format.space_before = Pt(8)
                first = par
            if i == len(lines) - 1:
                par.paragraph_format.space_after = Pt(2)
            par.add_run(line if line else " ")
        if caption:
            self.lst_n += 1
            self._caption("Listagem", self.lst_n, caption)
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(6)
        return first

    # ---- índices
    def toc(self, title="Índice"):
        self.h(1, title, in_toc=False)
        par = self.doc.add_paragraph()
        _field(par, ' TOC \\o "1-3" \\h \\z \\u ',
               placeholder="O índice é gerado automaticamente. No Word, selecione este texto e "
                           "prima F9 (ou clique com o botão direito \u2192 Atualizar campo).")
        for r in par.runs:
            r.font.size = Pt(10)
            r.font.color.rgb = SOFT

    def list_of(self, title, label):
        self.h(1, title)
        par = self.doc.add_paragraph()
        _field(par, f' TOC \\h \\z \\c "{label}" ',
               placeholder=f"Selecione este texto e prima F9 no Word para gerar a lista de "
                           f"{label.lower()}s.")
        for r in par.runs:
            r.font.size = Pt(10)
            r.font.color.rgb = SOFT

    def hrule(self):
        p = self.doc.add_paragraph()
        pbdr = _el("w:pBdr")
        pbdr.append(_el("w:bottom", **{"w:val": "single", "w:sz": "6",
                                       "w:space": "1", "w:color": RULE}))
        p._p.get_or_add_pPr().append(pbdr)
        p.paragraph_format.space_after = Pt(10)

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.doc.save(path)
        return path
